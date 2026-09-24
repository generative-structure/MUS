#!/usr/bin/env python3
"""Build MUS_Risk_Aligned_Allocation_Calculator.xlsx (openpyxl, formulas only)."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
import math

wb = Workbook()
ws = wb.active; ws.title = "Calculator"
F = "Arial"
bold = Font(name=F, bold=True); norm = Font(name=F); blue = Font(name=F, color="0000FF"); title = Font(name=F, bold=True, size=14)
green = Font(name=F, color="008000")
yellow = PatternFill("solid", fgColor="FFFF00"); grey = PatternFill("solid", fgColor="EEEEEE"); head = PatternFill("solid", fgColor="D9E1F2")
thin = Side(style="thin", color="999999"); box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="center")

ws["A1"] = "MUS Risk-Aligned Allocation Calculator"; ws["A1"].font = title
ws["A2"] = ("Allocation aid: distributes a probabilistic sample size you have already chosen across amount bands, after certainty items are removed. "
            "It does not determine sample size. Blue cells on yellow are inputs; everything else is a formula."); ws["A2"].font = norm
ws.merge_cells("A2:S2"); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[2].height = 32

ws["A4"] = "Mode (Basic or Advanced)"; ws["B4"] = "Advanced"
ws["C4"] = "Basic: dollar share and occurrence only (severity and cost assumed equal across bands). Advanced: adds severity, and cost when the budget basis is Cost."
ws["A5"] = "Budget basis (Count or Cost)"; ws["B5"] = "Count"
ws["C5"] = ("Count: B6 is a fixed number of transactions; with the count fixed, the variance-minimizing allocation does not depend on review cost, so cost is not used. "
            "Cost: B6 is an expected review budget in cost units (column G); cost enters the allocation and the expected counts are outputs.")
ws["A6"] = "Planned sample size (Count) or expected review budget (Cost)"; ws["B6"] = 400
ws["C6"] = "Chosen by the auditor from the engagement's own sample-size method; this workbook only allocates it."
ws["A7"] = "Reference band (1-4) for column Q"; ws["B7"] = 2
ws["C7"] = "Column Q shows the occurrence at which each band's risk multiplier would equal the reference band's. MUS is recovered only when every band's multiplier is equal."
ws["A8"] = "Check: dollar shares sum to 100%"; ws["B8"] = '=IF(ABS(D14-1)<0.0005,"OK","Dollar shares must sum to 100%")'
for r in (4, 5, 6, 7): ws[f"B{r}"].font = blue; ws[f"B{r}"].fill = yellow; ws[f"B{r}"].border = box
for r in (4, 5, 6, 7, 8): ws[f"A{r}"].font = bold; ws[f"C{r}"].font = Font(name=F, italic=True, size=9)
ws["B8"].font = bold
dv = DataValidation(type="list", formula1='"Basic,Advanced"', allow_blank=False); ws.add_data_validation(dv); dv.add("B4")
dvb = DataValidation(type="list", formula1='"Count,Cost"', allow_blank=False); ws.add_data_validation(dvb); dvb.add("B5")
dv2 = DataValidation(type="whole", operator="between", formula1="1", formula2="4"); ws.add_data_validation(dv2); dv2.add("B7")

headers = ["Band", "Label / amount range", "Number of transactions", "% of remaining dollars (D_h)",
           "Expected error occurrence % (p_h)", "Expected RMS % error when erroneous (sqrt of s_h)",
           "Expected review cost per item (Cost basis only)", "Severity used (s_h)", "Cost used (c_h)",
           "Risk multiplier sqrt(p*s/c), relative to reference band", "Unnormalized allocation weight D_h x multiplier",
           "MUS share (dollar-proportional)", "Suggested sample % (risk-aligned)", "MUS sample count",
           "Suggested sample count", "Difference (suggested minus MUS)", "Occurrence matching the reference band's risk multiplier",
           "Occurrence at which this band's share equals its MUS share (other bands fixed)",
           "Suggested coverage (% of band's transactions)", "MUS coverage (% of band's transactions)",
           "helper: suggested exact", "helper: floor", "helper: remainder", "helper: rank",
           "helper: MUS exact", "helper: floor", "helper: remainder", "helper: rank"]
cols = "ABCDEFGHIJKLMNOPQRST"
allcols = list(cols) + ["U", "V", "W", "X", "Y", "Z", "AA", "AB"]
for c, h in zip(allcols, headers):
    cell = ws[f"{c}9"]; cell.value = h; cell.font = bold; cell.fill = head; cell.alignment = wrap; cell.border = box
ws.row_dimensions[9].height = 90

# example inputs (hypothetical planning assumptions; match manuscript Table 7)
example = [
    (1, "Below $1,000", 6000, 0.10, 0.010, math.sqrt(0.30), 1),
    (2, "$1,000 to $9,999", 2500, 0.30, 0.015, math.sqrt(0.15), 1),
    (3, "$10,000 to $99,999", 800, 0.40, 0.020, math.sqrt(0.06), 1),
    (4, "$100,000 to certainty threshold", 120, 0.20, 0.025, math.sqrt(0.02), 1),
]
REF = "INDEX($E$10:$E$13,$B$7)*INDEX($H$10:$H$13,$B$7)/INDEX($I$10:$I$13,$B$7)"
for i, (b, lab, n, d, p, rms, cost) in enumerate(example):
    r = 10 + i
    ws[f"A{r}"] = b; ws[f"B{r}"] = lab; ws[f"C{r}"] = n; ws[f"D{r}"] = d; ws[f"E{r}"] = p; ws[f"F{r}"] = rms; ws[f"G{r}"] = cost
    for c in "BCDEFG":
        ws[f"{c}{r}"].font = blue; ws[f"{c}{r}"].fill = yellow
    ws[f"H{r}"] = f'=IF(AND($B$4="Advanced",F{r}>0),F{r}^2,1)'
    ws[f"I{r}"] = f'=IF(AND($B$4="Advanced",$B$5="Cost",G{r}>0),G{r},1)'
    ws[f"J{r}"] = f'=IF(OR(D{r}=0,E{r}=0),0,SQRT(E{r}*H{r}/I{r})/SQRT({REF}))'
    ws[f"K{r}"] = f'=D{r}*J{r}'
    ws[f"L{r}"] = f'=IF($D$14=0,0,D{r}/$D$14)'
    ws[f"M{r}"] = f'=IF($K$14=0,0,K{r}/$K$14)'
    # largest-remainder rounding to the target totals in B17 (suggested) and B18 (MUS)
    ws[f"U{r}"] = f'=M{r}*$B$17'; ws[f"V{r}"] = f'=INT(U{r})'; ws[f"W{r}"] = f'=U{r}-V{r}'
    ws[f"X{r}"] = f'=IF(U{r}=0,99,1+COUNTIF($W$10:$W$13,">"&W{r})+COUNTIF($W$9:W{r-1},W{r}))'
    ws[f"Y{r}"] = f'=L{r}*$B$18'; ws[f"Z{r}"] = f'=INT(Y{r})'; ws[f"AA{r}"] = f'=Y{r}-Z{r}'
    ws[f"AB{r}"] = f'=IF(Y{r}=0,99,1+COUNTIF($AA$10:$AA$13,">"&AA{r})+COUNTIF($AA$9:AA{r-1},AA{r}))'
    # Count basis: largest remainder to the planned total. Cost basis: round each band down, so cost <= budget.
    ws[f"N{r}"] = f'=IF($B$5="Cost",Z{r},Z{r}+IF(AB{r}<=$B$18-SUM($Z$10:$Z$13),1,0))'
    ws[f"O{r}"] = f'=IF($B$5="Cost",V{r},V{r}+IF(X{r}<=$B$17-SUM($V$10:$V$13),1,0))'
    ws[f"P{r}"] = f'=O{r}-N{r}'
    ws[f"Q{r}"] = f'=IF(OR(D{r}=0,H{r}=0),"",{REF}*I{r}/H{r})'
    ws[f"R{r}"] = f'=IF(OR(D{r}=0,H{r}=0,L{r}>=1),"",I{r}/H{r}*(SQRT({REF})*($K$14-K{r})/$D$14/(1-L{r}))^2)'
    ws[f"S{r}"] = f'=IF(C{r}=0,"",O{r}/C{r})'
    ws[f"T{r}"] = f'=IF(C{r}=0,"",N{r}/C{r})'
    for c in allcols: ws[f"{c}{r}"].border = box
    for c in "DELMQRST": ws[f"{c}{r}"].number_format = "0.00%"
    ws[f"F{r}"].number_format = "0.0%"; ws[f"C{r}"].number_format = "#,##0"; ws[f"G{r}"].number_format = "0.00"
    ws[f"H{r}"].number_format = "0.0000"; ws[f"I{r}"].number_format = "0.00"; ws[f"J{r}"].number_format = "0.000"; ws[f"K{r}"].number_format = "0.00000"
    for c in "NOP": ws[f"{c}{r}"].number_format = "0;(0);-"
    for c in ("U", "W", "Y", "AA"): ws[f"{c}{r}"].number_format = "0.0000"
    for c in ("U", "V", "W", "X", "Y", "Z", "AA", "AB"): ws[f"{c}{r}"].font = Font(name=F, color="808080", size=9)
ws["A14"] = "Total"; ws["A14"].font = bold
for c in "CDKLMNOP":
    ws[f"{c}14"] = f"=SUM({c}10:{c}13)"; ws[f"{c}14"].font = bold; ws[f"{c}14"].fill = grey; ws[f"{c}14"].border = box
ws["C14"].number_format = "#,##0"; ws["D14"].number_format = "0.00%"; ws["K14"].number_format = "0.00000"
ws["L14"].number_format = "0.00%"; ws["M14"].number_format = "0.00%"
for c in "NOP": ws[f"{c}14"].number_format = "0;(0);-"

ws["A16"] = "Totals and review cost"; ws["A16"].font = bold
ws["A17"] = "Total expected count, suggested allocation"
ws["B17"] = '=IF($B$5="Cost",IF(SUMPRODUCT(M10:M13,I10:I13)=0,0,$B$6/SUMPRODUCT(M10:M13,I10:I13)),$B$6)'
ws["A18"] = "Total expected count, MUS allocation"
ws["B18"] = '=IF($B$5="Cost",IF(SUMPRODUCT(L10:L13,I10:I13)=0,0,$B$6/SUMPRODUCT(L10:L13,I10:I13)),$B$6)'
ws["A21"] = "Budget check"
ws["B21"] = '=IF($B$5="Cost",IF(AND(B19<=$B$6+0.000001,B20<=$B$6+0.000001),"Within budget","OVER BUDGET"),"Count basis")'
ws["A19"] = "Expected review cost, suggested allocation"; ws["B19"] = "=SUMPRODUCT(O10:O13,I10:I13)"
ws["A20"] = "Expected review cost, MUS allocation"; ws["B20"] = "=SUMPRODUCT(N10:N13,I10:I13)"
for r in (17, 18, 19, 20, 21):
    ws[f"A{r}"].font = norm; ws[f"B{r}"].font = bold; ws[f"B{r}"].fill = grey; ws[f"B{r}"].border = box
    ws[f"B{r}"].number_format = "#,##0.00" if r >= 19 else "#,##0.0"
ws["C17"] = ("Count basis: equals B6, and band counts are allocated by largest-remainder rounding, so they sum exactly to it. "
             "Cost basis: the continuous expected count is the budget divided by the expected cost per selection under that allocation; "
             "each band's count is then rounded down, so the rounded allocation never costs more than the budget (B21 confirms this), "
             "and the unspent remainder is less than one selection's cost per band.")
ws["C17"].font = Font(name=F, italic=True, size=9); ws.merge_cells("C17:T18"); ws["C17"].alignment = Alignment(wrap_text=True, vertical="top")

ws["A22"] = "What MUS assumes"; ws["A22"].font = bold
ws["A23"] = ("The global MUS allocation is recovered only when all band multipliers sqrt(p_h x s_h / c_h) are equal. Column Q shows, for each band, the occurrence rate at which its "
             "multiplier would equal the reference band's; MUS is the allocation only if every band meets its column-Q value at once. Column R shows, holding the other bands "
             "at their entered values, the occurrence at which this band's own share would equal its MUS share; the other bands' shares still differ from MUS unless their multipliers also match.")
ws.merge_cells("A23:T23"); ws["A23"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[23].height = 44
ws["A25"] = "Formulas"; ws["A25"].font = bold
ws["A26"] = ("Suggested share_h = D_h x sqrt(p_h x s_h / c_h) / sum over bands of the same quantity. MUS share_h = D_h. Basic mode sets s_h = 1 and c_h = 1 for every band (equal severity and cost assumed). "
             "Count basis sets c_h = 1 (with a fixed count, cost does not change the variance-minimizing allocation). Cost basis uses c_h and fixes the expected review budget.")
ws["A27"] = ("s_h is entered as the RMS percentage error among erroneous transactions and squared inside the workbook. p_h x s_h stands for the band's dollar-weighted mean squared taint; "
             "entering it as a product of occurrence and severity assumes both are homogeneous within the band. Label the source of every input in the Notes sheet.")
ws["A28"] = ("Source: Equation (11) and Supplemental Sections S1.1 and S1.6 of 'Assumed, Not Measured'. Column R: p_h = (c_h / s_h) x [sum over other bands of D_j sqrt(p_j s_j / c_j) / (1 - D_h)]^2. "
             "Example values are hypothetical planning assumptions and match Table 7 of the manuscript.")
for r in (26, 27, 28):
    ws.merge_cells(f"A{r}:T{r}"); ws[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="top"); ws.row_dimensions[r].height = 30; ws[f"A{r}"].font = Font(name=F, size=9)
ws["A30"] = "Legend"; ws["A30"].font = bold
ws["A31"] = "Blue text on yellow fill = input you edit. Black = formula (do not overwrite). Grey columns U to AB are rounding helpers. Rows 10 to 13 example values are illustrative, not measured data."
ws.merge_cells("A31:T31"); ws["A31"].font = Font(name=F, size=9)
ws["D10"].comment = Comment("Share of the remaining recorded dollars in the band after certainty items are removed. Shares must sum to 100%.", "Calculator")
ws["E10"].comment = Comment("Anticipated probability that a transaction in this band contains an error. Planning judgment unless measured; see Notes.", "Calculator")
ws["F10"].comment = Comment("Anticipated root-mean-square percentage misstatement among erroneous transactions in the band (e.g. 0.30 means 30%). Squared inside the workbook to give s_h = E[taint^2 | error]. Used only in Advanced mode.", "Calculator")
ws["G10"].comment = Comment("Relative review cost per transaction. Used only in Advanced mode with the Cost budget basis; leave 1 if costs are equal.", "Calculator")
widths = {"A": 8, "B": 26, "C": 13, "D": 13, "E": 13, "F": 15, "G": 13, "H": 11, "I": 9, "J": 14, "K": 14, "L": 12, "M": 13, "N": 10, "O": 10, "P": 12, "Q": 15, "R": 16, "S": 13, "T": 13}
for c, w in widths.items(): ws.column_dimensions[c].width = w
for c in ("U", "V", "W", "X", "Y", "Z", "AA", "AB"): ws.column_dimensions[c].width = 9
ws.freeze_panes = "A10"

# ---------------- Sensitivity sheet ----------------
sv = wb.create_sheet("Sensitivity")
sv["A1"] = "Sensitivity: one band's occurrence rate"; sv["A1"].font = title
sv["A2"] = ("Vary the occurrence rate of one band while holding the other bands at their Calculator values. Multipliers apply to that band's p_h only; "
            "scaling every band by the same factor would leave the shares unchanged."); sv.merge_cells("A2:H2"); sv["A2"].alignment = Alignment(wrap_text=True, vertical="top"); sv.row_dimensions[2].height = 32
sv["A4"] = "Band to vary (1-4)"; sv["B4"] = 4; sv["B4"].font = blue; sv["B4"].fill = yellow; sv["B4"].border = box; sv["A4"].font = bold
dv3 = DataValidation(type="whole", operator="between", formula1="1", formula2="4"); sv.add_data_validation(dv3); dv3.add("B4")
sv["A5"] = "Base occurrence of that band"; sv["B5"] = "=INDEX(Calculator!$E$10:$E$13,$B$4)"; sv["B5"].number_format = "0.00%"; sv["A5"].font = bold; sv["B5"].font = green
sv["A6"] = "Base unnormalized weight of that band"; sv["B6"] = "=INDEX(Calculator!$K$10:$K$13,$B$4)"; sv["B6"].number_format = "0.00000"; sv["A6"].font = bold; sv["B6"].font = green
sv["A7"] = "Sum of other bands' weights"; sv["B7"] = "=Calculator!$K$14-$B$6"; sv["B7"].number_format = "0.00000"; sv["A7"].font = bold; sv["B7"].font = green
hdr = ["Multiplier on p_h", "Occurrence used", "Share: band 1", "Share: band 2", "Share: band 3", "Share: band 4", "Count: varied band", "MUS count: varied band"]
for c, h in zip("ABCDEFGH", hdr):
    cell = sv[f"{c}9"]; cell.value = h; cell.font = bold; cell.fill = head; cell.alignment = wrap; cell.border = box
sv.row_dimensions[9].height = 36
mults = [0.25, 0.5, 0.75, 1.0, 1.5, 2.0, 3.0, 5.0]
for i, mlt in enumerate(mults):
    r = 10 + i
    sv[f"A{r}"] = mlt; sv[f"A{r}"].font = blue; sv[f"A{r}"].fill = yellow; sv[f"A{r}"].number_format = "0.00"
    sv[f"B{r}"] = f"=$B$5*A{r}"; sv[f"B{r}"].number_format = "0.00%"
    # varied band's new weight = base weight * sqrt(multiplier); denominator = others + new weight
    for j, c in enumerate("CDEF", start=1):
        sv[f"{c}{r}"] = (f"=IF($B$7+$B$6*SQRT(A{r})=0,0,IF($B$4={j},$B$6*SQRT(A{r}),INDEX(Calculator!$K$10:$K$13,{j}))/($B$7+$B$6*SQRT(A{r})))")
        sv[f"{c}{r}"].number_format = "0.0%"
    sv[f"G{r}"] = (f'=IF(Calculator!$B$5="Cost",ROUNDDOWN(INDEX(C{r}:F{r},$B$4)*Calculator!$B$6/(C{r}*Calculator!$I$10+D{r}*Calculator!$I$11+E{r}*Calculator!$I$12+F{r}*Calculator!$I$13),0),ROUND(INDEX(C{r}:F{r},$B$4)*Calculator!$B$6,0))')
    sv[f"G{r}"].number_format = "0"
    sv[f"H{r}"] = "=INDEX(Calculator!$N$10:$N$13,$B$4)"; sv[f"H{r}"].number_format = "0"
    for c in "ABCDEFGH": sv[f"{c}{r}"].border = box
sv["A19"] = "Reading: the row with multiplier 1.00 reproduces the Calculator sheet's shares. Rows above and below show how much the varied band's share and count move if its occurrence rate is lower or higher than assumed. The varied band's count is rounded on its own, so it can differ by one from the largest-remainder count on the Calculator sheet."
sv.merge_cells("A19:H19"); sv["A19"].alignment = Alignment(wrap_text=True, vertical="top"); sv.row_dimensions[19].height = 30; sv["A19"].font = Font(name=F, size=9)
for c, w in zip("ABCDEFGH", (16, 14, 13, 13, 13, 13, 16, 18)): sv.column_dimensions[c].width = w

# ---------------- Notes sheet ----------------
nt = wb.create_sheet("Notes")
notes = [
    ("Purpose", "Allocates a chosen probabilistic sample across 3 or 4 amount bands in proportion to the residual misstatement risk the auditor anticipates in each band. Companion to 'Assumed, Not Measured: What Monetary-Unit Sampling Assumes About Where Error Lives'."),
    ("Not a sample-size tool", "The planned sample size or review budget (Calculator!B6) comes from the engagement's own sample-size determination. This workbook only distributes it across bands."),
    ("Certainty items", "Remove individually material, high-consequence and policy-required items before using the workbook. The dollar shares D_h refer to the remaining population."),
    ("Formula", "Suggested share_h = D_h x sqrt(p_h x s_h / c_h), normalized to sum to one. D_h = share of remaining dollars; p_h = probability a transaction in the band is erroneous; s_h = expected squared percentage misstatement among erroneous transactions (entered as its square root, the RMS percentage error); c_h = relative review cost per transaction."),
    ("Budget basis: Count", "B6 is a fixed number of transactions. With the count fixed, the variance-minimizing allocation is share_h proportional to D_h x sqrt(p_h x s_h) and does not depend on review cost, so c_h is set to 1. Use this basis when review costs are equal."),
    ("Budget basis: Cost", "B6 is an expected review budget in the cost units of column G. The allocation D_h x sqrt(p_h x s_h / c_h) is then the variance-minimizing one for that budget, and the expected counts (B17, B18) are outputs. Use this basis when review costs differ across bands."),
    ("Basic mode", "Sets s_h = 1 and c_h = 1 for every band, so share_h is proportional to D_h x sqrt(p_h). This assumes equal percentage severity and equal review cost across bands; state that assumption when you use it."),
    ("Advanced mode", "Uses the entered RMS percentage error (squared) as s_h, and under the Cost basis the entered cost as c_h. Blank or zero severity falls back to 1; blank or zero cost falls back to 1."),
    ("MUS as the special case", "The global MUS allocation is recovered only when all band multipliers sqrt(p_h x s_h / c_h) are equal; the suggested share then equals the MUS share D_h in every band."),
    ("Column Q", "Occurrence matching the reference band's risk multiplier: the rate at which this band's sqrt(p_h x s_h / c_h) would equal the reference band's. MUS is the allocation only if every band meets its column-Q value at the same time. Changing one band's occurrence to its column-Q value does not by itself make MUS the risk-aligned allocation."),
    ("Column R", "Holding the other bands at their entered values, the occurrence at which this band's suggested share equals its MUS share: p_h = (c_h / s_h) x [sum over other bands of D_j sqrt(p_j s_j / c_j) / (1 - D_h)]^2. The other bands' shares still differ from MUS unless their multipliers also match."),
    ("Within-band homogeneity", "The general band quantity is q_h = (1/A_h) x sum over items in the band of m_i / a_i, the band's dollar-weighted mean squared taint. Entering it as p_h x s_h assumes occurrence and severity are homogeneous within the band; with heterogeneity, enter p_h and s_h so that their product equals the band's dollar-weighted mean of occurrence times squared taint."),
    ("Where the inputs come from", "Prior audits of the same population, pilot or reference samples, re-performance, adjudicated historical results, or explicit planning assumptions. Inputs that are judgments rather than measurements should be labeled as such and varied in the Sensitivity sheet."),
    ("Within-band selection", "The rule assumes dollar-proportional selection within each band (the usual MUS mechanics) and changes only how the sample is divided between bands."),
    ("Rounding", "Count basis: band counts are allocated from the planned total by largest-remainder rounding and sum exactly to it. Cost basis: each band's continuous expected count is rounded down, so the expected review cost of the rounded allocation (B19, B20) never exceeds the budget; B21 reports the check. The unspent remainder is less than one selection's cost in each band and can be assigned by hand if it fits."),
    ("Example values", "The example bands (rows 10-13) are hypothetical planning inputs and match Table 7 of the manuscript. They are not measured data."),
    ("Verification", "The companion Python script check_calculator.py reproduces the example allocation independently and compares it with the workbook's recalculated values, in both budget bases."),
]
nt["A1"] = "Notes and assumptions"; nt["A1"].font = title
for i, (k, v) in enumerate(notes):
    r = 3 + i
    nt[f"A{r}"] = k; nt[f"A{r}"].font = bold; nt[f"A{r}"].alignment = Alignment(vertical="top")
    nt[f"B{r}"] = v; nt[f"B{r}"].font = norm; nt[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    nt.row_dimensions[r].height = 52
nt.column_dimensions["A"].width = 26; nt.column_dimensions["B"].width = 110

wb.calculation.fullCalcOnLoad = True
wb.save("MUS_Risk_Aligned_Allocation_Calculator.xlsx")
print("saved")
