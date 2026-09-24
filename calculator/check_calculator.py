#!/usr/bin/env python3
"""Independent check of MUS_Risk_Aligned_Allocation_Calculator.xlsx after LibreOffice/Excel recalculation.

Usage: python3 check_calculator.py [recalculated.xlsx ...]. Each workbook is checked in whatever mode and budget
basis it was saved with; the manuscript Table 7 values are asserted when the workbook holds the example in Count basis.
"""
import math, sys
from openpyxl import load_workbook

ok = 0
def check(a, b, msg, tol=1e-9):
    global ok
    assert abs(a - b) <= tol, f"{msg}: sheet {a} vs python {b}"; ok += 1


def ok_add():
    global ok
    ok += 1


def largest_remainder(shares, total):
    exact = [s * total for s in shares]; fl = [math.floor(x) for x in exact]
    order = sorted(range(len(exact)), key=lambda i: (-(exact[i] - fl[i]), i))
    for i in [j for j in order if exact[j] > 0][: total - sum(fl)]: fl[i] += 1
    return fl


def check_book(path):
    wb = load_workbook(path, data_only=True); ws = wb["Calculator"]; sv = wb["Sensitivity"]
    mode, basis, B, ref = ws["B4"].value, ws["B5"].value, ws["B6"].value, ws["B7"].value
    bands = []
    for r in range(10, 14):
        D, p, rms, cost = (ws[f"{c}{r}"].value for c in "DEFG")
        s = rms ** 2 if (mode == "Advanced" and rms) else 1.0
        c = cost if (mode == "Advanced" and basis == "Cost" and cost) else 1.0
        bands.append((D, p, s, c))
    xr = math.sqrt(bands[ref - 1][1] * bands[ref - 1][2] / bands[ref - 1][3])
    x = [math.sqrt(p * s / c) if D and p else 0 for D, p, s, c in bands]
    Dt = sum(b[0] for b in bands)
    w = [D * xi for (D, *_), xi in zip(bands, x)]
    share = [v / sum(w) for v in w]; mus = [b[0] / Dt for b in bands]; cost = [b[3] for b in bands]
    T = B / sum(sh * c for sh, c in zip(share, cost)) if basis == "Cost" else B
    Tm = B / sum(sh * c for sh, c in zip(mus, cost)) if basis == "Cost" else B
    if basis == "Cost":   # each band rounded down: expected cost can never exceed the budget
        cnt, mcnt = [math.floor(sh * T + 1e-9) for sh in share], [math.floor(sh * Tm + 1e-9) for sh in mus]
    else:
        cnt, mcnt = largest_remainder(share, T), largest_remainder(mus, Tm)
        check(ws["O14"].value, T, "suggested counts sum to total"); check(ws["N14"].value, Tm, "MUS counts sum to total")
    check(ws["B17"].value, T, "total suggested count", 1e-9); check(ws["B18"].value, Tm, "total MUS count", 1e-9)
    check(ws["B19"].value, sum(n * c for n, c in zip(cnt, cost)), "suggested review cost", 1e-9)
    check(ws["B20"].value, sum(n * c for n, c in zip(mcnt, cost)), "MUS review cost", 1e-9)
    if basis == "Cost":
        assert ws["B19"].value <= B + 1e-9 and ws["B20"].value <= B + 1e-9, "rounded allocation exceeds budget"; ok_add()
        assert ws["B21"].value == "Within budget", ws["B21"].value; ok_add()
    for i, r in enumerate(range(10, 14)):
        D, p, s, c = bands[i]
        check(ws[f"J{r}"].value, x[i] / xr, f"multiplier band {i+1}")
        check(ws[f"M{r}"].value, share[i], f"suggested share band {i+1}")
        check(ws[f"L{r}"].value, mus[i], f"MUS share band {i+1}")
        check(ws[f"O{r}"].value, cnt[i], f"suggested count band {i+1}")
        check(ws[f"N{r}"].value, mcnt[i], f"MUS count band {i+1}")
        check(ws[f"Q{r}"].value, xr ** 2 * c / s, f"occurrence matching reference multiplier band {i+1}")
        S = sum(bands[j][0] / Dt * x[j] for j in range(4) if j != i)
        check(ws[f"R{r}"].value, c / s * (S / (1 - mus[i])) ** 2, f"own-share occurrence band {i+1}", 1e-12)
        # at the column-R occurrence, this band's share equals its MUS share
        xi = math.sqrt(ws[f"R{r}"].value * s / c)
        check(mus[i] * xi / (mus[i] * xi + S), mus[i], f"column R reproduces MUS share band {i+1}", 1e-12)
    b = sv["B4"].value
    for r in range(10, 18):
        m = sv[f"A{r}"].value; wv = w[b - 1] * math.sqrt(m); den = sum(w) - w[b - 1] + wv
        sh = [(wv if j == b - 1 else w[j]) / den for j in range(4)]
        for j, col in enumerate("CDEF"):
            check(sv[f"{col}{r}"].value, sh[j], f"sensitivity row {r} band {j+1}")
        if basis == "Cost":
            expect = math.floor(sh[b - 1] * B / sum(s_ * c_ for s_, c_ in zip(sh, cost)) + 1e-9)
        else:
            expect = round(sh[b - 1] * B)
        check(sv[f"G{r}"].value, expect, f"sensitivity count row {r}")
    if mode == "Advanced" and basis == "Count" and B == 400:
        for i, (es, ec, ej, er) in enumerate(zip([0.144, 0.374, 0.364, 0.118], [57, 150, 146, 47],
                                                 [0.0075, 0.015, 0.0375, 0.1125], [0.00436, 0.00771, 0.02707, 0.08801])):
            check(round(share[i], 3), es, f"Table 7 share band {i+1}"); check(cnt[i], ec, f"Table 7 count band {i+1}")
            check(ws[f"Q{10+i}"].value, ej, f"Table 7 column Q band {i+1}"); check(round(ws[f"R{10+i}"].value, 5), er, f"column R band {i+1}")
    print(f"{path}: mode={mode} basis={basis} B6={B} shares={[round(v,4) for v in share]} counts={cnt} MUS={mcnt} totals={T}/{Tm}")


# MUS special case: equal multipliers give shares equal to D_h
D = [0.1, 0.3, 0.4, 0.2]; we = [d * math.sqrt(0.02) for d in D]
assert all(abs(v / sum(we) - d) < 1e-12 for v, d in zip(we, D)); ok += 1
for p in (sys.argv[1:] or ["MUS_Risk_Aligned_Allocation_Calculator.xlsx"]):
    check_book(p)
print(f"ALL {ok} CALCULATOR CHECKS PASSED")
